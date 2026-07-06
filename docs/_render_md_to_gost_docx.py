# -*- coding: utf-8 -*-
"""
Мини-конвертер Markdown → ГОСТ .docx (для глав диссертации; держит Word синхронным с .md).
Поддержка: заголовки # ## ###, абзацы с **жирным** и `кодом`, списки «- », таблицы «|...|»,
цитаты «> », горизонтальная черта «---».
Рендерит docs/Глава_1_*ПРОЕКТ_2026-06-16.md → одноимённый .docx.
Запуск: PYTHONUTF8=1 python docs/_render_md_to_gost_docx.py
"""
import glob
import os
import re

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = glob.glob(os.path.join(HERE, "Глава_1_*ПРОЕКТ_2026-06-16.md"))[0]
OUT = SRC[:-3] + ".docx"
FONT = "Times New Roman"
CONTENT_CM = 16.5  # A4 21 - (3.0 + 1.5)


def _font(run, size=14, bold=False, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


def _add_inline(p, text, size=14, base_bold=False):
    """Разбор **жирного**; backticks убираются."""
    text = text.replace("`", "")
    for k, part in enumerate(text.split("**")):
        if part == "":
            continue
        _font(p.add_run(part), size=size, bold=(base_bold or (k % 2 == 1)))


def _strip_md(s):
    return s.replace("**", "").replace("`", "")


def title(doc, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.space_after = Pt(8); pf.line_spacing = 1.5
    _font(p.add_run(_strip_md(text)), size=15, bold=True)


def heading(doc, text, size=14):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(10); pf.space_after = Pt(5); pf.line_spacing = 1.5; pf.keep_with_next = True
    _font(p.add_run(_strip_md(text)), size=size, bold=True)


def note(doc, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.line_spacing = 1.3; pf.space_after = Pt(5); pf.left_indent = Cm(0.5)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_inline(p, text, size=12)
    for r in p.runs:
        r.font.italic = True


def body(doc, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_after = Pt(6); pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_inline(p, text, size=14)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_after = Pt(3); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_inline(p, text, size=14)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcpr.append(shd)


def emit_table(doc, block):
    rows = []
    for ln in block:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    # отбросить строку-разделитель |---|
    data = [r for r in rows if not all(set(c) <= set("-: ") and c for c in r)]
    if not data:
        return
    ncols = len(data[0])
    u = CONTENT_CM / (ncols + 1.0)
    widths = [2 * u] + [u] * (ncols - 1)
    t = doc.add_table(rows=0, cols=ncols); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for ri, r in enumerate(data):
        cells = t.add_row().cells
        for ci in range(ncols):
            val = r[ci] if ci < len(r) else ""
            if ri == 0:
                _shade(cells[ci], "D9E2F3")
            cells[ci].vertical_alignment = 1
            par = cells[ci].paragraphs[0]; pf = par.paragraph_format
            pf.line_spacing = 1.0; pf.space_after = Pt(1); pf.space_before = Pt(1)
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            _add_inline(par, val, size=10, base_bold=(ri == 0))
    for ci in range(ncols):
        for r in t.rows:
            r.cells[ci].width = Cm(widths[ci])
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)


def add_page_numbers(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(); _font(run, size=11)
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    b = OxmlElement("w:instrText"); b.set(qn("xml:space"), "preserve"); b.text = "PAGE"
    c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), "end")
    run._element.append(a); run._element.append(b); run._element.append(c)


# ---------- сборка ----------
doc = Document()
nrm = doc.styles["Normal"]; nrm.font.name = FONT; nrm.font.size = Pt(14)
nrm.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); nrm.element.rPr.rFonts.set(qn("w:cs"), FONT)
s = doc.sections[0]
s.page_width = Cm(21.0); s.page_height = Cm(29.7)
s.left_margin = Cm(3.0); s.right_margin = Cm(1.5); s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
add_page_numbers(s)

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    raw = lines[i]; ln = raw.strip()
    if not ln or ln == "---":
        i += 1; continue
    if ln.startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i]); i += 1
        emit_table(doc, block); continue
    if ln.startswith("### "):
        heading(doc, ln[4:], size=13); i += 1; continue
    if ln.startswith("## "):
        heading(doc, ln[3:], size=14); i += 1; continue
    if ln.startswith("# "):
        title(doc, ln[2:]); i += 1; continue
    if ln.startswith("> "):
        note(doc, ln[2:]); i += 1; continue
    if ln.startswith("- "):
        bullet(doc, ln[2:]); i += 1; continue
    body(doc, ln); i += 1

# fix zoom + save
zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")
doc.save(OUT)
print("OK:", os.path.basename(OUT))
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
print("bytes:", os.path.getsize(OUT))
