# -*- coding: utf-8 -*-
"""
Сборка Word-записки «Обзор готовых программных решений и разграничение
научной новизны» по результатам prior-art-сканирования (США/ЕС/КНР/РФ).

Источник содержания: docs/novelty/prior_art_map.md (UPDATE 2026-06-15 + добор
Китай/Европа) — этот скрипт лишь оформляет наблюдения в .docx по ГОСТ-стилю.

Запуск (из корня репозитория):  PYTHONUTF8=1 python docs/novelty/_build_prior_art_review_docx.py
Печать в консоль — только ASCII (cp1251-консоль не выводит UTF-8).
"""
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Обзор_готовых_решений_2026-06-16.docx")

FONT = "Times New Roman"


# ---------- низкоуровневые помощники ----------
def _set_run_font(run, size=14, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def body(doc, text, size=14, justify=True, indent=True, bold=False, italic=False,
         space_after=6, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    if indent:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def heading(doc, text, size=14, before=12, after=6, center=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=True)
    return p


def bullet(doc, text, size=14):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    _set_run_font(run, size=size)
    return p


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _cell(cell, text, size=10.5, bold=False, fill=None, align="left", color=None):
    if fill:
        _shade(cell, fill)
    cell.vertical_alignment = 1  # center
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(1)
    pf.space_before = Pt(1)
    pf.first_line_indent = Cm(0)
    pf.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, color=color)


def table(doc, headers, rows, widths_cm, font=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for j, h in enumerate(headers):
        _cell(hdr[j], h, size=font, bold=True, fill="D9E2F3", align="center")
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            al = "center" if j == 0 and False else ("center" if val in ("Да", "Нет", "ДА", "—", "Частично") else "left")
            color = None
            if val == "ДА":
                color = RGBColor(0x1F, 0x6F, 0x3C)
            _cell(cells[j], val, size=font, align=al, color=color,
                  bold=(val in ("ДА",)))
    # ширины колонок (нужно проставлять каждой ячейке)
    for j, w in enumerate(widths_cm):
        for r in t.rows:
            r.cells[j].width = Cm(w)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(8)
    pf.space_after = Pt(3)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_with_next = True
    run = p.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=True)


def add_page_numbers(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _set_run_font(run, size=11)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._element.append(fld1); run._element.append(instr); run._element.append(fld2)


# ---------- документ ----------
doc = Document()

# базовый стиль
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(14)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.element.rPr.rFonts.set(qn("w:cs"), FONT)

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(1.5)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
add_page_numbers(sec)

# --- титул ---
heading(doc, "Обзор готовых программных решений по тематике диссертации\n"
             "и разграничение научной новизны", size=15, before=0, after=6, center=True)
p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
_set_run_font(p.add_run("Аналитическая записка по результатам поиска аналогов "
                        "(коммерческие, открытые и академические решения; США, ЕС, КНР, РФ)"),
              size=12, italic=True)
p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
_set_run_font(p.add_run("Специальность 2.4.2 «Электротехнические комплексы и системы». "
                        "Объект: outrunner SPM PMSM 1–5 кВт (БПЛА), NdFeB N42SH, сталь M270-35A"),
              size=12)
p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
_set_run_font(p.add_run("Дата: 16 июня 2026 г.   Исполнитель: Рубцов С.А."), size=12)

# --- 1. Краткое резюме ---
heading(doc, "1. Краткое резюме")
body(doc, "Цель обзора — убедиться, что разрабатываемый решатель не повторяет уже "
          "существующих продуктов и научных результатов, и зафиксировать защитимое "
          "«белое поле» (незанятую комбинацию возможностей).")
body(doc, "Главный вывод. Ни один из найденных инструментов — коммерческий, открытый "
          "или академический, в США, Европейском союзе, Китае и России — не объединяет "
          "в ОДНОЙ расчётной модели две вещи одновременно: (а) расчёт открытой "
          "(неограниченной) области через ИСТИННЫЙ метод граничных элементов (BEM), без "
          "искусственного «обрезания» расчётной области, и (б) нелинейное температурно-"
          "зависимое НЕОБРАТИМОЕ размагничивание постоянных магнитов с поэлементной картой "
          "риска. Именно эта интеграция и является защитимым ядром работы.")
body(doc, "Существенная оговорка. Каждая из двух «половин» по отдельности — зрелый "
          "опубликованный задел: открытый расчёт FEM/BEM (например, Kielhorn и др., 2017) "
          "и модели размагничивания с картой риска (ANSYS Maxwell, JMAG; работы Ruoho 2011, "
          "Kern 2023). Поэтому новизна носит ИНТЕГРАТИВНЫЙ характер — защищается их сборка, "
          "а не отдельные компоненты.")
body(doc, "Две корректировки формулировок новизны, выявленные обзором: (1) научную новизну "
          "по adjoint-оптимизации (положение 3) следует СУЗИТЬ — этот подход для магнитных "
          "машин занят широко (группа Gangl/RICAM, корейская и иные школы); (2) «сертифицированный "
          "суррогат» (положение 4) следует ПЕРЕПОЗИЦИОНИРОВАТЬ — сертифицированные редуцированные "
          "модели с апостериорными оценками ошибки — зрелый европейский задел (в т.ч. для "
          "вращающихся машин), поэтому конформное предсказание (conformal prediction) надо "
          "подавать как дополнение, а не как единственную «сертификацию».")

# --- 2. Метод и охват ---
heading(doc, "2. Метод и охват, ограничения достоверности")
body(doc, "Использованы: (1) автоматизированное «глубокое исследование» с веерным поиском по "
          "пяти направлениям, загрузкой первоисточников и состязательной перекрёстной проверкой "
          "фактов (подтверждено 25 из 25 ключевых утверждений голосованием 3 из 3); (2) прицельный "
          "поиск по остаточным пробелам и по китайским/европейским источникам. Приоритет отдавался "
          "первоисточникам: документации производителей, рецензируемым статьям, диссертациям и "
          "авторефератам.")
body(doc, "Ограничение. Поисковый бэкенд преимущественно англоязычный. Международные журналы "
          "(IEEE, COMPEL, IET, MDPI, Springer, ScienceDirect), где публикуется конкурентная "
          "работа КНР и ЕС, охвачены. Однако китайские базы (CNKI, Wanfang) и русские диссертации "
          "(dissercat, eLibrary) на национальных языках охвачены НЕполно. «Не найдено» не равно "
          "«не существует»: перед защитой целесообразен ручной добор по CNKI и eLibrary.")

# --- 3. Матрица ---
heading(doc, "3. Матрица «инструмент × возможность»")
caption(doc, "Таблица 1 — Сопоставление по ключевой оси: умеет ли инструмент нелинейное "
             "необратимое размагничивание и каким способом он считает открытую область")
table(
    doc,
    headers=["Инструмент", "Нелинейный необратимый demag\n(колено / T / карта риска)",
             "Метод открытой границы", "Истинный BEM-экстерьер\n+ demag в одной модели"],
    rows=[
        ["Altair Flux / FluxMotor", "Да", "«Infinite box» — трансформация координат (по докам Altair, «отлична от BEM»)", "Нет"],
        ["ANSYS Motor-CAD", "Да (колено по откл. ≥10 %)", "Balloon (через ANSYS Maxwell)", "Нет"],
        ["Siemens Simcenter MAGNET / Motorsolve", "Да (3D-карты «hot-spots»)", "Меш-усечение (воздух / дальняя граница)", "Нет"],
        ["COMSOL AC/DC", "Да (в FEM-части)", "Гибрид FEM-BEM есть, но BEM = скалярный No-Currents", "Нет (вендор: «BEM не для нелинейных»)"],
        ["ANSYS Maxwell", "Да (поэлементно, поле Demag_Coef, recoil)", "Balloon", "Нет"],
        ["JMAG-Designer", "Да («коэффициент размагничивания»)", "Меш-усечение", "Нет"],
        ["FEMM (+ pyFEMM / SyR-e)", "Частично", "IABC / преобразование Кельвина (усечение)", "Нет"],
        ["NGSolve / ngbem", "Нет (нет модели demag)", "Симметричный FEM/BEM, но только скаляр (Лаплас)", "Нет"],
        ["Kielhorn 2017 (академ., ЕС)", "Нет (магнит линеен, M=const)", "Истинный симметричный FEM/BEM (вектор-A)", "Нет"],
        ["Ступаков 2016 (академ., РФ)", "Нет (магнит линеен)", "FEM/BEM + быстрый мультипольный (скаляр)", "Нет"],
        ["Субдомен-FE гибрид (КНР, 2024)", "Да", "Субдомен-аналитика + FE (НЕ BEM)", "Нет"],
        ["Наш решатель", "Да", "Истинный FEM/BEM (Стеклов–Пуанкаре)", "ДА"],
    ],
    widths_cm=[4.6, 4.0, 4.4, 3.5],
)
body(doc, "Вывод по таблице 1: размагничивание умеют практически все; открытую область все "
          "промышленные пакеты закрывают УСЕЧЕНИЕМ (infinite box, balloon, преобразование Кельвина), "
          "а не граничными элементами. Истинный BEM-экстерьер ВМЕСТЕ с нелинейным размагничиванием "
          "не реализован ни в одном инструменте.", space_before=4)

# --- 4. По категориям и географиям ---
heading(doc, "4. Наблюдения по категориям и географиям")

heading(doc, "4.1. Коммерческие пакеты (США/ЕС/Япония)", size=13, before=8, after=4)
body(doc, "ANSYS Maxwell и Motor-CAD, Altair Flux/FluxMotor, Siemens Simcenter MAGNET/Motorsolve, "
          "JMAG, COMSOL — все имеют развитое нелинейное необратимое размагничивание (колено на "
          "внутренней кривой, recoil, температурная зависимость, 3D-карты). Открытую область они "
          "закрывают мешевым усечением. COMSOL единственный имеет общий гибрид FEM-BEM, но его BEM — "
          "скалярная формулировка «без токов», и сам производитель указывает, что BEM неприменим к "
          "нелинейным/неоднородным материалам, то есть размагничивание там жить не может.")

heading(doc, "4.2. Открытые и академические решения", size=13, before=8, after=4)
body(doc, "Истинный открытый FEM/BEM для электрических машин — это прежде всего симметричная "
          "вектор-A/Неделек схема Kielhorn–Rüberg–Zechner (2017) и изогеометрическая FEM/BEM "
          "ТУ Дармштадта; в обеих магнит ЛИНЕЕН (M=const), размагничивания нет. Все найденные "
          "решатели размагничивания — замкнутый/пошаговый FEM без открытой границы (Ruoho 2011 и др.). "
          "NGSolve предоставляет симметричную связь FEM/BEM, но только для скалярной задачи Лапласа.")

heading(doc, "4.3. Россия", size=13, before=8, after=4)
body(doc, "Ближайший отечественный аналог по связке FEM/BEM — диссертация И.М. Ступакова (НГТУ, 2016). "
          "Однако это специальность 05.13.18 (математическое моделирование), а не 2.4.2; используется "
          "скалярный потенциал и быстрый мультипольный BEM для нелинейной СТАЛИ, магниты ЛИНЕЙНЫ, "
          "приложение — ускорители частиц. Прицельный поиск по запросам «программный комплекс + "
          "синхронный двигатель с постоянными магнитами + необратимое размагничивание» дал только "
          "работы по управлению и диагностике. Отечественного открытого решателя FEM/BEM с нелинейным "
          "размагничиванием и картой риска не найдено.")

heading(doc, "4.4. Китай и Азия", size=13, before=8, after=4)
body(doc, "Объём работ по размагничиванию PMSM в КНР огромен, но расчёты идут в ЗАМКНУТОМ/пошаговом "
          "FEM (в т.ч. с электромагнитно-тепловой связью). «Гибрид» в этих работах означает связку "
          "субдомен-аналитики с FE или Максвелл-Фурье с FE, но НЕ граничные элементы для внешней "
          "области; BEM применяют к расчёту силы линейных магнитов. Индигенного открытого решателя "
          "FEM/BEM с нелинейным размагничиванием не обнаружено; проектирование PMSM опирается на "
          "цепочки коммерческих пакетов (Motor-CAD + optiSLang + ANSYS).")

heading(doc, "4.5. Европа", size=13, before=8, after=4)
body(doc, "Европа — источник почти всех «соседей» по двум рискованным положениям. Adjoint/"
          "топологическая оптимизация магнитных машин (в т.ч. с ограничением по размагничиванию) — "
          "группа Gangl/Krenn (RICAM, Австрия) и смежные. Сертифицированные редуцированные модели "
          "(строгие апостериорные оценки ошибки) для электромагнетики и даже для вращающихся машин — "
          "школа Hesthaven–Rozza–Stamm и Haasdonk. Это поднимает планку для положения 4.")

# --- 5. Разграничение новизны ---
heading(doc, "5. Разграничение по пяти положениям научной новизны")
caption(doc, "Таблица 2 — Статус каждого положения, ближайший задел и рекомендуемая формулировка")
table(
    doc,
    headers=["Положение", "Статус", "Ближайший prior-art", "Как заявлять"],
    rows=[
        ["(1) Связка FEM/BEM (открытая область)",
         "Инфраструктура (не самостоятельная новизна)",
         "Kielhorn 2017 (ЕС); Ступаков 2016 (РФ); Salgado–Selgas 2008",
         "Как корректную верифицированную архитектуру — носитель физики; не как науч. новизну"],
        ["(2) Нелин. B(H,T) + колено + карта необратимого размагничивания на FEM-стороне ОТКРЫТОГО решателя",
         "БЕЛОЕ ПОЛЕ — держится (США/ЕС/КНР/РФ)",
         "Модель: Ruoho 2011, Kern 2023; карта риска: ANSYS Maxwell, JMAG, замкнутый FEM КНР",
         "Заявлять ИНТЕГРАЦИЮ (физика внутри открытого FEM/BEM), а не саму модель магнита"],
        ["(3) Adjoint-чувствительности сквозь колено",
         "Занято широко → узкий остаток",
         "Gangl/Krenn (Австрия); корейская школа (level-set + CDSA); Putek; кит./франц. TO",
         "Узко: adjoint сквозь открытое нелин. колено demag-FEM/BEM + конформная маржа; не «adjoint машин» вообще"],
        ["(4) Сертифицированный суррогат (FOM→ROM→NN + conformal + trust-region)",
         "Высокий риск — обе половины заняты",
         "Апостериорные ROM: Hesthaven–Rozza–Stamm, Haasdonk (вкл. вращ. машины); conformal: UQNO, RESS 2025",
         "Conformal — как distribution-free дополнение к апостериорным оценкам (для NN-уровня); фиксировать поздно"],
        ["(5) Интегрированный программный комплекс",
         "Аналога связки нет ни в РФ, ни в КНР",
         "Коммерч. цепочки (Motor-CAD + optiSLang + ANSYS)",
         "Носитель интеграции и практической значимости; не самостоятельная науч. новизна"],
    ],
    widths_cm=[3.8, 3.0, 4.4, 5.3],
    font=10.0,
)

# --- 6. Красные флаги ---
heading(doc, "6. «Красные флаги» (что НЕ заявлять как новое в одиночку)")
bullet(doc, "Открытый расчёт FEM/BEM — занят (Kielhorn 2017; Ступаков 2016).")
bullet(doc, "Модель размагничивания B(H,T) и карта риска — заняты (ANSYS Maxwell, JMAG; Ruoho, Kern; замкнутый FEM КНР).")
bullet(doc, "Магнитотепловое размагничивание — занято в замкнутом FEM (КНР).")
bullet(doc, "Adjoint-оптимизация магнитных машин с учётом размагничивания — занято глобально (Австрия, Корея, Франция, КНР).")
bullet(doc, "Сертифицированные ROM с апостериорными оценками, в т.ч. для вращающихся машин — заняты (европейская школа).")
bullet(doc, "Conformal prediction на суррогате и в петле проектирования — заняты (UQNO; RESS 2025).")

# --- 7. Белое поле ---
heading(doc, "7. Защитимое «белое поле»")
body(doc, "Незанятой остаётся именно ИНТЕГРАЦИЯ положения (2): нелинейный анизотропный "
          "температурно-зависимый магнит B(H,T) с детекцией колена и трекингом необратимого "
          "размагничивания, встроенный на FEM-сторону открытого (FEM/BEM) решателя с поэлементной "
          "картой риска, для outrunner SPM PMSM, с верификацией от первооснов (три аналитических "
          "бенчмарка + метод многообразных решений). Узкие остатки положений (3) и (4) "
          "формулировать осторожно (см. таблицу 2); положение (5) — как носитель интеграции и "
          "практической значимости.")

# --- 8. Рекомендации ---
heading(doc, "8. Рекомендации")
bullet(doc, "Ядро (положение 2) — развивать уверенно: оно подтверждено как незанятое на четырёх географиях.")
bullet(doc, "Положение (3) — сузить формулировку до «adjoint сквозь открытое нелинейное колено + конформная маржа».")
bullet(doc, "Положение (4) — переформулировать: conformal как distribution-free дополнение к апостериорным ROM, а не замена.")
bullet(doc, "Перед защитой — ручной добор по CNKI/Wanfang (через коллегу/перевод) и eLibrary для закрытия национально-язычного пробела.")
bullet(doc, "Синхронизировать документ «Научная_новизна_2_4_2.docx» с этими корректировками через гейт научного руководителя.")

# --- 9. Источники ---
heading(doc, "9. Ключевые источники (сверить полный текст перед защитой)")
src = [
    "Kielhorn L., Rüberg T., Zechner J. Simulation of electrical machines — a FEM-BEM coupling scheme // COMPEL 36(5), 2017. arXiv:1610.05472.",
    "Ступаков И.М. Разработка алгоритмов решения задач магнитостатики с использованием МГЭ: дис. … канд. техн. наук, НГТУ, 2016 (ВАК 05.13.18).",
    "Ruoho S. Modeling Demagnetization of Sintered NdFeB Magnet Material … : PhD, Aalto, 2011.",
    "Kern A., Leuning N., Hameyer K. Semi-physical demagnetization model … // AIP Advances 13:025105, 2023.",
    "ANSYS Maxwell, Help (v242): Irreversible Demagnetization Due to Temperature Change (поле Demag_Coef).",
    "JMAG-International: thermal demagnetization of IPM motor (коэффициент размагничивания).",
    "Altair Flux, User Guide: Infinite box transformation (метод открытой границы, отличный от BEM).",
    "Siemens Simcenter MAGNET/Motorsolve: 3D demagnetization (blogs.sw.siemens.com).",
    "Krenn N., Gangl P. Multi-material topology optimization … considering demagnetization. arXiv:2404.12188 (2024); Robust TO. arXiv:2504.05070; Electro-thermal TO. arXiv:2507.14759.",
    "Topology optimization of rotor poles … level set + continuum design sensitivity analysis (adjoint).",
    "Hesthaven J., Rozza G., Stamm B. Certified Reduced Basis Methods …; Certified RB для EFIE (SIAM SISC, doi:10.1137/110848268).",
    "Dihlmann M., Haasdonk B. Certified PDE-constrained parameter optimization …; Model Order Reduction for Rotating Electrical Machines (Springer).",
    "Partovizadeh, Schöps, Loukrezis. Fourier-enhanced reduced-order surrogate … для UQ в проектировании машин. arXiv:2412.06485 (MC-UQ, без conformal).",
    "Uncertainty quantification of surrogate models using conformal prediction. arXiv:2408.09881; Sequential surrogate modeling … with conformal inference (RESS, 2025).",
    "Demagnetization of PMSM based on Subdomain-Finite Element Hybrid Method (Springer, 2024) — «гибрид» = субдомен+FE, не BEM.",
    "Подробный реестр с цитатами и статусами: docs/novelty/prior_art_map.md (UPDATE 2026-06-15 + добор Китай/Европа).",
]
for i, s in enumerate(src, 1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(2)
    pf.left_indent = Cm(0.75)
    pf.first_line_indent = Cm(-0.75)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_run_font(p.add_run("%d. %s" % (i, s)), size=11)

# python-docx кладёт <w:zoom> без обязательного по схеме w:percent — добавляем
_settings = doc.settings.element
_zoom = _settings.find(qn("w:zoom"))
if _zoom is not None and _zoom.get(qn("w:percent")) is None:
    _zoom.set(qn("w:percent"), "100")

doc.save(OUT)
print("OK: saved docx")
print("paragraphs:", len(doc.paragraphs))
print("path-bytes:", os.path.getsize(OUT))
